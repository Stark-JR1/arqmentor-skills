from __future__ import annotations
from pathlib import Path
import sys, json, re
from openpyxl import load_workbook

ERROR_TOKENS = ("#REF!", "#NAME?", "#VALUE!", "#DIV/0!")

def inspect(path: Path):
    wb = load_workbook(path, data_only=False, read_only=False, keep_links=True)
    report = {"file": str(path), "sheets": [], "formula_errors": []}
    for ws in wb.worksheets:
        info = {"name": ws.title, "max_row": ws.max_row, "max_column": ws.max_column, "state": ws.sheet_state, "formula_count": 0}
        for row in ws.iter_rows():
            for c in row:
                v=c.value
                if isinstance(v, str) and v.startswith('='):
                    info["formula_count"] += 1
                    for token in ERROR_TOKENS:
                        if token in v:
                            report["formula_errors"].append({"sheet": ws.title, "cell": c.coordinate, "formula": v, "token": token})
        report["sheets"].append(info)
    return report

if __name__ == '__main__':
    if len(sys.argv) != 2:
        raise SystemExit('Uso: python workbook_healthcheck.py arquivo.xlsx')
    p=Path(sys.argv[1])
    print(json.dumps(inspect(p), ensure_ascii=False, indent=2))
